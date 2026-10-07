"""Dựng dữ liệu bước 2 cho Cá Chef: ghi các tab ra data/*.csv và data/ca_chef_data.xlsx.

Lịch tháng: 2 = đang rộ/ngon nhất, 1 = có hàng, 0 = trái mùa/hiếm.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from recipes import RECIPES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
GO_CSV = os.path.join(ROOT, "research", "go_catalog_2026-10-07.csv")
GO_DATE = "2026-10-07"


def months(have=(), peak=(), all_year=False):
    """Trả về list 12 giá trị cho T1..T12."""
    row = [1 if all_year else 0] * 12
    for m in have:
        row[m - 1] = max(row[m - 1], 1)
    for m in peak:
        row[m - 1] = 2
    return row


def span(a, b):
    """Khoảng tháng a..b, cho phép vắt qua năm (vd 11..4)."""
    return list(range(a, b + 1)) if a <= b else list(range(a, 13)) + list(range(1, b + 1))


# (mã, tên, nhóm, vùng, lịch tháng, nơi mua, tin cậy, ghi chú, tên sản phẩm GO! khớp)
INGREDIENTS = [
    # --- Ninh Thuận ---
    ("nho_nt", "Nho Ninh Thuận", "trái cây", "Ninh Thuận", months(all_year=True, peak=[8, 9]), "chợ/vựa", "A", "2–3 vụ/năm: 4–5, 8–9, Tết; ngọt nhất 8–9", None),
    ("tao_xanh", "Táo xanh Phan Rang", "trái cây", "Ninh Thuận", months(all_year=True, peak=[12, 1, 2]), "chợ/vựa", "B", "Vụ tập trung dịp Tết", None),
    ("toi_pr", "Tỏi Phan Rang", "gia vị", "Ninh Thuận", months(all_year=True, peak=[1, 2, 3]), "chợ/vựa", "B", "Thu hoạch quanh Tết; có nguồn ghi 4–5", None),
    ("hanh_tim", "Hành tím Ninh Thuận", "gia vị", "Ninh Thuận", months(all_year=True), "chợ", "B", "Chu kỳ ~45 ngày, thu rải nhiều đợt", None),
    ("mang_tay", "Măng tây xanh", "rau củ", "Ninh Thuận", months(all_year=True), "chợ/vựa", "A", "2 vụ trồng, có quanh năm", None),
    ("thit_de", "Thịt dê", "thịt", "Ninh Thuận", months(all_year=True), "chợ/quán", "A", "", None),
    ("thit_cuu", "Thịt cừu", "thịt", "Ninh Thuận", months(all_year=True), "chợ/quán", "A", "", None),
    ("tom_hum", "Tôm hùm (Vĩnh Hy, Bình Ba)", "hải sản", "Ninh Thuận/Khánh Hòa", months(all_year=True), "vựa", "B", "Nuôi lồng", None),
    ("rong_sun", "Rong sụn Ninh Hải", "hải sản", "Ninh Thuận", months(all_year=True), "chợ", "C", "", None),
    ("ca_com", "Cá cơm tươi", "hải sản", "Ninh Thuận", months(have=span(4, 9), peak=[7, 8]), "chợ", "A", "Vụ cá Nam 4–9", None),
    ("ca_nuc", "Cá nục", "hải sản", "Ninh Thuận", months(have=span(4, 9), peak=[7, 8]), "chợ/GO!", "A", "Vụ cá Nam 4–9", "Cá nục hấp GO! 250g"),
    ("muc_tuoi", "Mực tươi", "hải sản", "Ninh Thuận", months(have=span(4, 9), peak=[7, 8]), "chợ/GO!", "A", "Vụ cá Nam 4–9", "Mực ống 300g (size 20-30 con/kg)"),
    ("muc_mot_nang", "Mực một nắng", "hải sản", "Ninh Thuận", months(all_year=True), "chợ/vựa", "B", "Hàng phơi, mùa biển động vẫn có", None),
    ("ca_mai", "Cá mai", "hải sản", "Ninh Thuận", months(have=span(4, 9)), "chợ", "C", "Chưa có nguồn về mùa, tạm theo vụ cá Nam", None),
    ("sua", "Sứa", "hải sản", "Ninh Thuận", months(have=span(3, 6)), "chợ", "B", "Tháng 2–5 âm lịch", None),
    ("ca_com_kho", "Cá cơm khô", "hàng khô", "Ninh Thuận", months(all_year=True), "chợ", "B", "", None),
    ("mam_ca_na", "Mắm/nước mắm Cà Ná", "gia vị", "Ninh Thuận", months(all_year=True), "chợ", "A", "", None),
    ("muoi_ca_na", "Muối Cà Ná", "gia vị", "Ninh Thuận", months(all_year=True), "chợ", "A", "Thu hoạch chính 4–7", None),
    # --- Tỉnh lân cận ---
    ("thanh_long", "Thanh long", "trái cây", "Bình Thuận", months(all_year=True, peak=[7, 8, 9]), "GO!/chợ", "A", "Chính vụ 6–12, trái vụ 1–5", "Thanh long (1.5-1.8kg)"),
    ("sau_rieng_ks", "Sầu riêng Khánh Sơn", "trái cây", "Khánh Hòa", months(have=[7, 8, 9], peak=[8]), "vựa", "B", "Chín muộn hơn miền Tây 3–4 tháng", None),
    ("sau_rieng_dl", "Sầu riêng Đắk Lắk", "trái cây", "Đắk Lắk", months(have=[7, 8, 9], peak=[7, 8]), "vựa/chợ", "A", "GO! online không bán sầu riêng tươi", None),
    ("mang_cut", "Măng cụt Bảo Lộc", "trái cây", "Lâm Đồng", months(have=[8, 9, 10], peak=[9, 10]), "chợ", "B", "", None),
    ("bo_booth", "Bơ booth", "trái cây", "Đắk Lắk", months(have=[8, 9, 10], peak=[8, 9]), "GO!/chợ", "B", "", "Bơ Booth 1-1.2kg"),
    ("bo_sap", "Bơ sáp", "trái cây", "Đắk Lắk/Lâm Đồng", months(have=span(3, 8), peak=[5, 6, 7]), "GO!/chợ", "B", "", "Bơ sáp - 3kg"),
    ("xoai_uc", "Xoài Úc Cam Lâm", "trái cây", "Khánh Hòa", months(have=span(3, 7) + span(9, 12)), "chợ/vựa", "B", "2 vụ; GO! online chỉ có xoài cát", None),
    ("chom_chom", "Chôm chôm Khánh Sơn", "trái cây", "Khánh Hòa", months(have=[6, 7, 8]), "chợ", "C", "", None),
    ("dua_hau", "Dưa hấu", "trái cây", "Gia Lai/miền Trung", months(all_year=True, peak=[12, 1, 2, 3, 4, 5, 6]), "GO!/chợ", "B", "Vụ Tết 12–2, vụ hè 3–6", "Dưa hấu ruột đỏ (2-2.5kg/trái)"),
    ("ca_ngu_dd", "Cá ngừ đại dương", "hải sản", "Phú Yên", months(have=span(1, 5), peak=[4]), "GO!/chợ", "A", "GO! có hàng đông lạnh quanh năm", "Cá ngừ đại dương đông lạnh Foodymart 500g"),
    # --- Đà Lạt ---
    ("bap_cai", "Bắp cải Đà Lạt", "rau củ", "Đà Lạt", months(all_year=True), "chợ", "A", "", None),
    ("ca_rot", "Cà rốt", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Cà rốt 800g-1kg"),
    ("sup_lo", "Súp lơ xanh", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Bông cải/súp lơ xanh 500-600g"),
    ("cai_thao", "Cải thảo", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Cải thảo (800g-1kg/bắp)"),
    ("xa_lach", "Xà lách", "rau củ", "Đà Lạt", months(all_year=True), "chợ", "A", "", None),
    ("khoai_tay", "Khoai tây", "rau củ", "Đà Lạt", months(all_year=True), "chợ", "A", "", None),
    ("su_su", "Su su", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Su su Đà Lạt 250g"),
    ("ca_chua", "Cà chua Đà Lạt", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Cà chua Đà Lạt 1 Kg"),
    ("nam", "Nấm trồng (bào ngư, kim châm...)", "rau củ", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Nấm bào ngư trắng 200g"),
    ("chanh_day", "Chanh dây", "trái cây", "Đà Lạt", months(all_year=True), "GO!/chợ", "A", "", "Chanh dây 1kg"),
    ("atiso", "Atiso tươi", "rau củ", "Đà Lạt", months(have=span(1, 4)), "chợ", "B", "", None),
    ("dau_tay", "Dâu tây", "trái cây", "Đà Lạt", months(have=span(11, 4), peak=[1, 2, 3]), "chợ/GO!", "A", "", None),
    ("man_dl", "Mận Đà Lạt", "trái cây", "Đà Lạt", months(have=[4, 5, 6]), "chợ", "C", "Chưa có nguồn", None),
    ("hong_gion", "Hồng giòn", "trái cây", "Đà Lạt", months(have=[9, 10, 11], peak=[10, 11]), "chợ", "A", "", None),
    ("hong_treo_gio", "Hồng treo gió", "trái cây", "Đà Lạt", months(have=[10, 11, 12]), "chợ", "A", "", None),
    ("mang_rung", "Măng rừng", "rau củ", "Đà Lạt", months(have=span(5, 9)), "chợ", "C", "Theo mùa mưa Tây Nguyên", None),
    # --- Nguyên liệu cơ bản (quanh năm) ---
    ("thit_heo", "Thịt heo", "thịt", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("long_heo", "Lòng heo", "thịt", "chung", months(all_year=True), "chợ", "A", "", None),
    ("ga_ta", "Gà ta", "thịt", "chung", months(all_year=True), "chợ", "A", "", None),
    ("trung", "Trứng gà/vịt", "trứng", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("tom_the", "Tôm thẻ", "hải sản", "chung", months(all_year=True), "GO!/chợ", "A", "", "Tôm thẻ 31-40 Minh Phú 300g"),
    ("ca_thu", "Cá thu", "hải sản", "chung", months(all_year=True), "GO!/chợ", "A", "", "Cá thu cắt khúc đông lạnh 500g"),
    ("bot_gao", "Bột gạo", "tinh bột", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("bun", "Bún / bánh hỏi / bánh canh", "tinh bột", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("banh_trang", "Bánh tráng", "tinh bột", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("dau_phong", "Đậu phộng", "gia vị", "chung", months(all_year=True), "GO!/chợ", "A", "", None),
    ("rau_thom", "Rau thơm các loại", "rau củ", "chung", months(all_year=True), "chợ", "A", "", None),
    ("khe_me", "Khế, me (vị chua)", "gia vị", "chung", months(all_year=True), "chợ", "A", "", None),
    ("sa_ot", "Sả, ớt, gừng", "gia vị", "chung", months(all_year=True), "chợ", "A", "", None),
]

# Quy tắc chấm điểm. Biến lấy từ Open-Meteo (daily).
# ap_dung_cho: <cột nhãn món>=<giá trị>; diem cộng/trừ vào điểm món.
RULES = [
    ("R01", "Nắng gắt", "temperature_2m_max", ">=", 34, "nhiet=mat", 3, "Ưu tiên món mát"),
    ("R02", "Nắng gắt", "temperature_2m_max", ">=", 34, "nhiet=nong", -2, "Hạn chế món nóng"),
    ("R03", "Nắng gắt", "temperature_2m_max", ">=", 34, "dau_mo=nhieu", -1, "Hạn chế món nhiều dầu"),
    ("R04", "Nóng vừa", "temperature_2m_max", "between", "31-34", "nhiet=mat", 1, ""),
    ("R05", "Trời mát", "temperature_2m_max", "<", 28, "nhiet=nong", 2, "Chủ yếu tháng 12–1"),
    ("R16", "Trời dịu", "temperature_2m_max", "between", "28-30", "nhiet=nong", 1, "Tháng 11–2 thường 28–30°C"),
    ("R06", "Có mưa", "precipitation_sum", ">=", 3, "nhiet=nong", 3, "Mưa ≥3mm/ngày"),
    ("R07", "Có mưa", "precipitation_sum", ">=", 3, "nhiet=mat", -1, ""),
    ("R08", "Mưa to", "precipitation_sum", ">=", 20, "do_nang=nang", 1, "Lẩu, món no bụng"),
    ("R09", "Khả năng mưa cao", "precipitation_probability_max", ">=", 60, "nhiet=nong", 1, "Chỉ áp dụng khi lượng mưa dự báo <3mm"),
    ("R10", "Gió mạnh", "wind_speed_10m_max", ">=", 25, "loai=canh", 1, "Gió km/h; năm 2025 gió mạnh nhất 34 km/h"),
    ("R11", "UV rất cao", "uv_index_max", ">=", 10, "nhiet=mat", 1, ""),
    ("R12", "Nguyên liệu chính đang rộ", "lich_thang", "=", 2, "nguyen_lieu_chinh", 3, "Lấy từ tab nguyen_lieu"),
    ("R13", "Nguyên liệu chính có hàng", "lich_thang", "=", 1, "nguyen_lieu_chinh", 1, ""),
    ("R14", "Nguyên liệu chính trái mùa", "lich_thang", "=", 0, "nguyen_lieu_chinh", -5, "Gần như loại khỏi gợi ý"),
    ("R15", "Vừa gợi ý gần đây", "so_ngay_tu_lan_cuoi", "<=", 3, "mon", -4, "Tránh lặp món"),
]

MON_AN_HEADER = ["ma_mon", "ten_mon", "loai", "nhiet", "do_nang", "dau_mo", "nguyen_lieu_chinh",
                 "thoi_tiet_hop", "khau_phan_goc", "thoi_gian_phut", "do_kho", "mo_ta_ngan",
                 "cach_lam", "do_pho_bien", "nguon", "trang_thai"]
MON_NL_HEADER = ["ma_mon", "ma_nguyen_lieu", "ten_hien_thi", "so_luong", "don_vi",
                 "kieu_tinh", "vai_tro", "ghi_chu"]

GUIDE = [
    ("Tab", "Nội dung"),
    ("nguyen_lieu", "Nguyên liệu + lịch 12 tháng. T1..T12: 2 = đang rộ, 1 = có hàng, 0 = trái mùa. tin_cay: A nhiều nguồn khớp, B 1 nguồn/lệch, C chưa có nguồn."),
    ("mon_an", "Danh sách món. nhiet: mat/am/nong. do_nang: nhe/nang. dau_mo: it/vua/nhieu. loai: mon_chinh/canh/goi/lau/an_vat/trang_mieng/do_uong. thoi_tiet_hop: nang/mua/moi. nguyen_lieu_chinh: nhiều mã cách nhau | = chỉ cần 1 mã đang có mùa. do_pho_bien: số kết quả tìm trên Cookpad VN (thô, để tham khảo). trang_thai: chua_nau_thu/da_nau_thu."),
    ("mon_nguyen_lieu", "Định lượng theo khẩu phần gốc. kieu_tinh: theo_nguoi (nhân thẳng) hoặc theo_noi (tăng chậm, cho gia vị/nước dùng). vai_tro: chinh/phu/gia_vi."),
    ("quy_tac", "Điểm món = tổng điểm các quy tắc khớp. App lấy 3 món điểm cao nhất, không trùng loại."),
    ("gia_go", f"Giá tham khảo từ sieuthi-go.vn ngày {GO_DATE} (danh mục online, có thể khác GO! Ninh Thuận)."),
]


def go_prices():
    catalog = {}
    with open(GO_CSV, newline="") as f:
        for name, price, avail, url in csv.reader(f):
            catalog.setdefault(name, (price, avail, url))
    rows = []
    for ing in INGREDIENTS:
        product = ing[8]
        if not product:
            continue
        if product not in catalog:
            sys.exit(f"Không thấy sản phẩm GO!: {product}")
        price, avail, url = catalog[product]
        rows.append([ing[0], product, int(float(price)), avail, url, GO_DATE])
    return rows


def recipe_rows():
    codes = {i[0] for i in INGREDIENTS}

    def check(field, ma):
        for c in field.split("|"):
            if c and c not in codes:
                sys.exit(f"{ma}: mã nguyên liệu lạ '{c}'")

    mon, mon_nl = [], []
    for r in RECIPES:
        check(r["chinh"], r["ma"])
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(r["cach_lam"], 1))
        mon.append([r["ma"], r["ten"], r["loai"], r["nhiet"], r["do_nang"], r["dau_mo"], r["chinh"],
                    r["thoi_tiet"], r["khau_phan"], r["phut"], r["do_kho"], r["mo_ta"], steps,
                    r["pho_bien"], r["nguon"], "chua_nau_thu"])
        for ma_nl, ten, sl, dv, kieu, vai in r["nguyen_lieu"]:
            check(ma_nl, r["ma"])
            mon_nl.append([r["ma"], ma_nl, ten, "" if sl is None else sl, dv,
                           "theo_nguoi" if kieu == "n" else "theo_noi", vai, ""])
    return mon, mon_nl


def build_tables():
    mon, mon_nl = recipe_rows()
    ing_rows = [[c, n, g, v, *lich, mua, tc, gc] for c, n, g, v, lich, mua, tc, gc, _ in INGREDIENTS]
    return {
        "huong_dan": [list(r) for r in GUIDE],
        "nguyen_lieu": [["ma", "ten", "nhom", "vung", *[f"T{i}" for i in range(1, 13)],
                         "noi_mua", "tin_cay", "ghi_chu"]] + ing_rows,
        "mon_an": [MON_AN_HEADER] + mon,
        "mon_nguyen_lieu": [MON_NL_HEADER] + mon_nl,
        "quy_tac": [["ma", "ten", "bien", "toan_tu", "nguong", "ap_dung_cho", "diem", "ghi_chu"]]
                   + [list(r) for r in RULES],
        "gia_go": [["ma_nguyen_lieu", "san_pham_go", "gia_vnd", "tinh_trang", "url", "ngay_lay"]] + go_prices(),
    }


def main():
    tables = build_tables()
    os.makedirs(DATA, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tables.items():
        with open(os.path.join(DATA, f"{name}.csv"), "w", newline="") as f:
            csv.writer(f).writerows(rows)
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFE699")
        ws.freeze_panes = "A2"
    wb.save(os.path.join(DATA, "ca_chef_data.xlsx"))
    for name, rows in tables.items():
        print(f"{name}: {len(rows) - 1} dòng")


if __name__ == "__main__":
    main()
